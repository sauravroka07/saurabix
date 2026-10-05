import os
import io
import base64

import streamlit as st
from groq import Groq
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook


# =========================================================
# APP SETTINGS
# =========================================================

APP_NAME = "Saurabix"

TEXT_MODEL = "openai/gpt-oss-120b"
VISION_MODEL = "qwen/qwen3.8-27b"

MAX_FILE_CHARACTERS = 80000


MODES = {
    "Chat": (
        "Answer questions clearly, naturally and accurately."
    ),

    "Explain Simply": (
        "Explain difficult topics using simple beginner-friendly "
        "language. Use examples when useful."
    ),

    "Code Helper": (
        "Help with programming, debugging and software development. "
        "Give correct working code and explain it clearly."
    ),

    "Quiz Maker": (
        "Create 5 useful multiple-choice questions about the requested "
        "topic. Use A, B, C and D and provide the answers at the end."
    ),

    "Nepali Helper": (
        "Reply in simple, natural and correct Nepali using Devanagari."
    ),

    "GK Helper": (
        "Answer general knowledge questions accurately and clearly. "
        "Add one useful interesting fact when appropriate."
    ),
}


BASE_PROMPT = (
    f"You are {APP_NAME}, a helpful AI assistant created by Saurav Roka, "
    "a BSc IT student from Nepal. "
    "Help users with studies, programming, general knowledge, "
    "problem solving and everyday questions. "
    "Give accurate, useful and natural answers. "
    "For difficult questions, reason carefully before answering. "
    "Do not make up information. "
    "If something is uncertain, clearly say so. "
    "Avoid unnecessary repetition."
)


# =========================================================
# PAGE SETTINGS
# =========================================================

st.set_page_config(
    page_title="Saurabix",
    page_icon="S",
    layout="wide",
    initial_sidebar_state="expanded",
)


# =========================================================
# DESIGN
# =========================================================

st.markdown(
    """
    <style>

    .stApp {
        background-color: #f8fafc;
    }

    .block-container {
        max-width: 1100px;
        padding-top: 2rem;
        padding-bottom: 6rem;
    }

    section[data-testid="stSidebar"] {
        background-color: #ffffff;
        border-right: 1px solid #dbe3ef;
    }

    section[data-testid="stSidebar"] * {
        color: #172033;
    }

    h1, h2, h3 {
        color: #172033 !important;
    }

    p {
        color: #475569;
    }

    div[data-baseweb="input"] {
        background-color: #ffffff !important;
        border: 1px solid #cbd5e1 !important;
        border-radius: 10px !important;
    }

    div[data-baseweb="input"]:focus-within {
        border-color: #3b82f6 !important;
        box-shadow: 0 0 0 3px rgba(59,130,246,.12) !important;
    }

    div[data-baseweb="input"] input {
        color: #111827 !important;
        background-color: #ffffff !important;
    }

    div[data-baseweb="input"] input::placeholder {
        color: #64748b !important;
    }

    div[data-baseweb="select"] > div {
        background-color: #ffffff !important;
        border-color: #cbd5e1 !important;
        border-radius: 10px !important;
    }

    .stButton > button {
        border-radius: 10px;
        border: 1px solid #cbd5e1;
        background-color: #ffffff;
        color: #172033;
        font-weight: 600;
        transition: .2s;
    }

    .stButton > button:hover {
        border-color: #3b82f6;
        color: #2563eb;
    }

    div[data-testid="stChatInput"] {
        position: sticky;
        bottom: 1rem;
        z-index: 999;
    }

    div[data-testid="stChatInput"] > div {
        background-color: #ffffff !important;
        border: 2px solid #cbd5e1 !important;
        border-radius: 18px !important;
        box-shadow: 0 8px 25px rgba(15,23,42,.10) !important;
    }

    div[data-testid="stChatInput"] > div:focus-within {
        border-color: #3b82f6 !important;
        box-shadow:
            0 0 0 4px rgba(59,130,246,.10),
            0 10px 30px rgba(15,23,42,.12) !important;
    }

    div[data-testid="stChatInput"] textarea {
        color: #111827 !important;
        background-color: transparent !important;
        font-size: 16px !important;
    }

    div[data-testid="stChatInput"] textarea::placeholder {
        color: #64748b !important;
    }

    div[data-testid="stChatMessage"] {
        padding-top: .5rem;
        padding-bottom: .5rem;
    }

    div[data-testid="stChatMessage"] p,
    div[data-testid="stChatMessage"] li {
        color: #172033 !important;
        line-height: 1.65;
    }

    div[data-testid="stAlert"] {
        border-radius: 14px;
    }

    hr {
        border-color: #e2e8f0;
    }

    #MainMenu {
        visibility: hidden;
    }

    footer {
        visibility: hidden;
    }

    </style>
    """,
    unsafe_allow_html=True,
)


# =========================================================
# API KEY
# =========================================================

api_key = os.getenv("GROQ_API_KEY", "")

try:
    api_key = api_key or st.secrets.get("GROQ_API_KEY", "")
except Exception:
    pass


# =========================================================
# SIDEBAR
# =========================================================

with st.sidebar:

    st.title("⚙️ Settings")

    st.caption("GROQ API")

    if not api_key:

        api_key = st.text_input(
            "Groq API key",
            type="password",
            placeholder="Paste your Groq API key",
        )

    st.caption("ASSISTANT MODE")

    mode = st.selectbox(
        "Choose a mode",
        list(MODES.keys())
    )

    st.divider()

    if st.button(
        "🗑️ New Chat",
        use_container_width=True
    ):

        st.session_state.messages = []
        st.session_state.uploaded_file = None
        st.session_state.file_context = ""
        st.session_state.file_name = ""
        st.session_state.file_type = ""
        st.session_state.image_data = None

        st.rerun()

    st.divider()

    st.caption("ABOUT")

    st.write("Saurabix")
    st.caption("Made by Saurav Roka")
    st.caption("Powered by Groq")


# =========================================================
# WELCOME SCREEN
# =========================================================

st.title("🔷 Saurabix")

st.caption(
    "Your AI study buddy for coding, learning, GK and more."
)

st.divider()


if not api_key:

    st.header("Welcome to Saurabix")

    st.write(
        "Your personal AI assistant is ready to help you "
        "learn, code and explore new ideas."
    )

    st.info(
        "🔑 Enter your Groq API key in the Settings panel "
        "to start chatting."
    )

    col1, col2, col3 = st.columns(3)

    with col1:

        st.subheader("💻 Coding")

        st.write(
            "Learn programming, understand code and solve problems."
        )

    with col2:

        st.subheader("📚 Study")

        st.write(
            "Understand difficult topics with simple explanations."
        )

    with col3:

        st.subheader("🧠 Knowledge")

        st.write(
            "Ask questions and explore general knowledge."
        )

    st.divider()

    st.caption("Saurabix • Made by Saurav Roka")

    st.stop()


# =========================================================
# GROQ CLIENT
# =========================================================

client = Groq(api_key=api_key)


# =========================================================
# SESSION STATE
# =========================================================

if "messages" not in st.session_state:
    st.session_state.messages = []

if "uploaded_file" not in st.session_state:
    st.session_state.uploaded_file = None

if "file_context" not in st.session_state:
    st.session_state.file_context = ""

if "file_name" not in st.session_state:
    st.session_state.file_name = ""

if "file_type" not in st.session_state:
    st.session_state.file_type = ""

if "image_data" not in st.session_state:
    st.session_state.image_data = None

if "file_warning" not in st.session_state:
    st.session_state.file_warning = ""


# =========================================================
# FILE READING
# =========================================================

def read_pdf(uploaded_file):

    reader = PdfReader(
        io.BytesIO(uploaded_file.getvalue())
    )

    pages = []

    for page_number, page in enumerate(reader.pages, start=1):

        text = page.extract_text()

        if text:
            pages.append(
                f"--- Page {page_number} ---\n{text}"
            )

    return "\n\n".join(pages), len(reader.pages)


def read_docx(uploaded_file):

    document = Document(
        io.BytesIO(uploaded_file.getvalue())
    )

    paragraphs = []

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            paragraphs.append(
                paragraph.text
            )

    # Also read tables
    for table in document.tables:

        for row in table.rows:

            values = []

            for cell in row.cells:

                if cell.text.strip():
                    values.append(cell.text.strip())

            if values:
                paragraphs.append(
                    " | ".join(values)
                )

    return "\n\n".join(paragraphs)


def read_txt(uploaded_file):

    data = uploaded_file.getvalue()

    try:

        return data.decode("utf-8")

    except UnicodeDecodeError:

        return data.decode("latin-1")


def read_xlsx(uploaded_file):

    workbook = load_workbook(
        filename=io.BytesIO(
            uploaded_file.getvalue()
        ),
        data_only=True
    )

    output = []

    for sheet in workbook.worksheets:

        output.append(
            f"--- Sheet: {sheet.title} ---"
        )

        for row in sheet.iter_rows(
            values_only=True
        ):

            values = []

            for value in row:

                if value is not None:

                    values.append(
                        str(value)
                    )

            if values:

                output.append(
                    " | ".join(values)
                )

    return "\n".join(output)


# =========================================================
# IMAGE PROCESSING
# =========================================================

def prepare_image(uploaded_file):

    image_bytes = uploaded_file.getvalue()

    encoded_image = base64.b64encode(
        image_bytes
    ).decode("utf-8")

    mime_type = uploaded_file.type

    return (
        f"data:{mime_type};base64,{encoded_image}"
    )


# =========================================================
# DISPLAY CHAT HISTORY
# =========================================================

for message in st.session_state.messages:

    with st.chat_message(message["role"]):

        st.markdown(
            message["content"]
        )


# =========================================================
# STARTER CARDS
# =========================================================

if not st.session_state.messages:

    st.subheader(
        "What can I help you with?"
    )

    st.caption(
        "Ask anything or try one of these examples."
    )

    col1, col2, col3 = st.columns(3)

    with col1:

        st.info(
            "💻 **Coding**\n\n"
            "Explain Python loops simply."
        )

    with col2:

        st.info(
            "📚 **Study**\n\n"
            "Explain how computer memory works."
        )

    with col3:

        st.info(
            "🧠 **Quiz**\n\n"
            "Quiz me on general knowledge."
        )


# =========================================================
# FILE UPLOAD
# =========================================================

st.divider()

st.subheader("📎 Upload a file")

st.caption(
    "Upload a PDF, Word document, text file, "
    "Excel spreadsheet or image."
)


uploaded_file = st.file_uploader(
    "Choose a file",
    type=[
        "pdf",
        "docx",
        "txt",
        "xlsx",
        "jpg",
        "jpeg",
        "png",
        "webp"
    ],
    label_visibility="collapsed"
)


# =========================================================
# PROCESS UPLOADED FILE
# =========================================================

if uploaded_file is not None:

    # Process only when a different file is selected
    if uploaded_file.name != st.session_state.file_name:

        st.session_state.file_context = ""
        st.session_state.image_data = None
        st.session_state.file_warning = ""

        file_name = uploaded_file.name.lower()

        try:

            # -------------------------------------------------
            # IMAGE
            # -------------------------------------------------

            if file_name.endswith(
                (".jpg", ".jpeg", ".png", ".webp")
            ):

                st.session_state.image_data = (
                    prepare_image(uploaded_file)
                )

                st.session_state.file_name = (
                    uploaded_file.name
                )

                st.session_state.file_type = (
                    "Image"
                )


            # -------------------------------------------------
            # PDF
            # -------------------------------------------------

            elif file_name.endswith(".pdf"):

                with st.spinner(
                    "Reading PDF..."
                ):

                    text, page_count = read_pdf(
                        uploaded_file
                    )

                st.session_state.file_context = text

                st.session_state.file_name = (
                    uploaded_file.name
                )

                st.session_state.file_type = (
                    f"PDF • {page_count} pages"
                )


            # -------------------------------------------------
            # DOCX
            # -------------------------------------------------

            elif file_name.endswith(".docx"):

                with st.spinner(
                    "Reading Word document..."
                ):

                    text = read_docx(
                        uploaded_file
                    )

                st.session_state.file_context = text

                st.session_state.file_name = (
                    uploaded_file.name
                )

                st.session_state.file_type = (
                    "Word document"
                )


            # -------------------------------------------------
            # TXT
            # -------------------------------------------------

            elif file_name.endswith(".txt"):

                with st.spinner(
                    "Reading text file..."
                ):

                    text = read_txt(
                        uploaded_file
                    )

                st.session_state.file_context = text

                st.session_state.file_name = (
                    uploaded_file.name
                )

                st.session_state.file_type = (
                    "Text file"
                )


            # -------------------------------------------------
            # XLSX
            # -------------------------------------------------

            elif file_name.endswith(".xlsx"):

                with st.spinner(
                    "Reading spreadsheet..."
                ):

                    text = read_xlsx(
                        uploaded_file
                    )

                st.session_state.file_context = text

                st.session_state.file_name = (
                    uploaded_file.name
                )

                st.session_state.file_type = (
                    "Excel spreadsheet"
                )


            # -------------------------------------------------
            # LIMIT LARGE DOCUMENTS
            # -------------------------------------------------

            if st.session_state.file_context:

                if (
                    len(
                        st.session_state.file_context
                    )
                    > MAX_FILE_CHARACTERS
                ):

                    st.session_state.file_context = (
                        st.session_state.file_context[
                            :MAX_FILE_CHARACTERS
                        ]
                    )

                    st.session_state.file_warning = (
                        "This file is very large. "
                        "Saurabix is currently using the "
                        "first part of the document."
                    )


        except Exception as error:

            st.session_state.file_name = ""
            st.session_state.file_type = ""
            st.session_state.file_context = ""
            st.session_state.image_data = None

            st.error(
                f"Could not read the file: {error}"
            )


# =========================================================
# FILE STATUS
# =========================================================

if st.session_state.file_name:

    if st.session_state.file_type == "Image":

        st.success(
            f"🖼️ {st.session_state.file_name} is ready."
        )

        st.image(
            st.session_state.image_data,
            caption=st.session_state.file_name,
            width=500
        )

        st.caption(
            "You can now ask Saurabix questions about this image."
        )

    else:

        st.success(
            f"📄 {st.session_state.file_name} is ready."
        )

        st.caption(
            f"Type: {st.session_state.file_type}"
        )

        if st.session_state.file_warning:

            st.warning(
                st.session_state.file_warning
            )

        if (
            st.session_state.file_type.startswith("PDF")
            and not st.session_state.file_context
        ):

            st.warning(
                "This PDF does not contain selectable text. "
                "It may be a scanned PDF. OCR support can "
                "be added later."
            )


# =========================================================
# TEXT AI FUNCTION
# =========================================================

def get_text_reply(
    history,
    system_prompt,
    file_context=""
):

    messages = [
        {
            "role": "system",
            "content": system_prompt
        }
    ]

    if file_context:

        messages.append(
            {
                "role": "system",
                "content": (
                    "The user uploaded a document. "
                    "Use this document as context when "
                    "answering questions about it.\n\n"
                    "DOCUMENT CONTENT:\n\n"
                    + file_context
                )
            }
        )

    for message in history:

        messages.append(
            {
                "role": message["role"],
                "content": message["content"]
            }
        )

    stream = client.chat.completions.create(
        model=TEXT_MODEL,
        messages=messages,
        temperature=0.5,
        reasoning_effort="medium",
        stream=True,
    )

    for chunk in stream:

        if not chunk.choices:
            continue

        delta = chunk.choices[0].delta

        if delta.content:

            yield delta.content


# =========================================================
# IMAGE AI FUNCTION
# =========================================================

def get_image_reply(
    history,
    system_prompt,
    image_data
):

    messages = [
        {
            "role": "system",
            "content": system_prompt
        }
    ]

    # Add previous text conversation
    for message in history[:-1]:

        messages.append(
            {
                "role": message["role"],
                "content": message["content"]
            }
        )

    # Current question + image
    current_question = history[-1]["content"]

    messages.append(
        {
            "role": "user",
            "content": [
                {
                    "type": "text",
                    "text": current_question
                },
                {
                    "type": "image_url",
                    "image_url": {
                        "url": image_data
                    }
                }
            ]
        }
    )

    response = client.chat.completions.create(
        model=VISION_MODEL,
        messages=messages,
        temperature=0.5,
        max_completion_tokens=2048,
        stream=False,
    )

    return response.choices[0].message.content


# =========================================================
# CHAT INPUT
# =========================================================

prompt = st.chat_input(
    "Message Saurabix..."
)


if prompt:

    st.session_state.messages.append(
        {
            "role": "user",
            "content": prompt
        }
    )

    with st.chat_message("user"):

        st.markdown(prompt)


    # -----------------------------------------------------
    # SYSTEM PROMPT
    # -----------------------------------------------------

    system_prompt = (
        BASE_PROMPT
        + "\n\nCurrent mode: "
        + MODES[mode]
    )


    # -----------------------------------------------------
    # DOCUMENT INSTRUCTIONS
    # -----------------------------------------------------

    if st.session_state.file_context:

        system_prompt += """

The user has uploaded a document.

When the question is about the uploaded document:

- Use the document as the main source.
- Answer directly and clearly.
- Do not invent information.
- If the answer cannot be found in the document,
  clearly say that.
- You may use general knowledge when the user asks
  for additional explanation.
"""


    # -----------------------------------------------------
    # IMAGE INSTRUCTIONS
    # -----------------------------------------------------

    if st.session_state.image_data:

        system_prompt += """

The user has uploaded an image.

Analyze the image carefully.

You can help with:
- Objects
- People
- Text visible in the image
- Screenshots
- Diagrams
- Charts
- Homework questions
- Programming screenshots
- Documents photographed by the user

Do not claim to see something that is not actually visible.
"""


    # -----------------------------------------------------
    # GENERATE RESPONSE
    # -----------------------------------------------------

    with st.chat_message("assistant"):

        try:

            # IMAGE RESPONSE
            if st.session_state.image_data:

                reply = get_image_reply(
                    st.session_state.messages,
                    system_prompt,
                    st.session_state.image_data
                )

                st.markdown(reply)


            # TEXT/DOCUMENT RESPONSE
            else:

                reply = st.write_stream(
                    get_text_reply(
                        st.session_state.messages,
                        system_prompt,
                        st.session_state.file_context
                    )
                )


        except Exception as error:

            reply = (
                "I couldn't process that request right now."
            )

            st.error(
                f"Something went wrong: {error}"
            )


    # -----------------------------------------------------
    # SAVE RESPONSE
    # -----------------------------------------------------

    st.session_state.messages.append(
        {
            "role": "assistant",
            "content": reply
        }
    )